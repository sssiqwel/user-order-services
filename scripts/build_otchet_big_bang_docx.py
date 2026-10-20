#!/usr/bin/env python3
import importlib.util
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.enum.text import WD_LINE_SPACING

HERE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("otchet_helpers", HERE / "build_otchet_docx.py")
h = importlib.util.module_from_spec(spec)
spec.loader.exec_module(h)

SRC = HERE.parent / "docs" / "otchet-big-bang.md"
OUT = HERE.parent / "docs" / "Otchet-big-bang.docx"
TABLE_WIDTH_CM = 16.5


def plain(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.replace("**", "").replace("`", "").rstrip()


def add_numbered_item(doc, number: str, text: str) -> None:
    p = doc.add_paragraph()
    p.add_run(f"{number}. {text}")
    h.style_paragraph(p, size=14, align="left", space_after=2, first_line=False)
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE


def add_bold_lead(doc, lead: str, rest: str) -> None:
    p = doc.add_paragraph()
    p.add_run(lead)
    if rest:
        p.add_run(" " + rest)
    h.style_paragraph(p, size=14, align="justify", space_after=6)
    p.runs[0].bold = True


def parse_row(line: str) -> list[str]:
    return [plain(cell.strip()) for cell in line.strip().strip("|").split("|")]


def build():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)

    styles = doc.styles["Normal"]
    styles.font.name = "Times New Roman"
    styles.font.size = Pt(14)
    styles._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    lines = SRC.read_text(encoding="utf-8").splitlines()
    in_header = True
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            h.add_code_block(doc, block)
            i += 1
            continue

        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.fullmatch(r"\|[\s\-:|]+\|", lines[i].strip()):
                    rows.append(parse_row(lines[i]))
                i += 1
            width = TABLE_WIDTH_CM / len(rows[0])
            h.add_table(doc, rows[0], rows[1:], col_widths=[width] * len(rows[0]))
            continue

        if not stripped:
            i += 1
            continue

        if stripped.startswith("# "):
            h.add_title(doc, plain(stripped[2:]))
        elif stripped.startswith("## ") and in_header:
            h.add_title(doc, plain(stripped[3:]))
            in_header = False
        elif stripped.startswith("## "):
            h.add_heading(doc, plain(stripped[3:]))
        elif stripped.startswith("### "):
            h.add_subheading(doc, plain(stripped[4:]))
        elif stripped.startswith("- "):
            h.add_bullet(doc, plain(stripped[2:]))
        elif m := re.match(r"(\d+)\.\s+(.*)", stripped):
            add_numbered_item(doc, m.group(1), plain(m.group(2)))
        elif m := re.match(r"\*\*([^*]+)\*\*\s*(.*)", stripped):
            lead, rest = m.group(1), plain(m.group(2))
            if lead.endswith(":") and len(rest) < 60 and i < 12:
                h.add_meta(doc, f"{lead} {rest}")
            else:
                add_bold_lead(doc, lead, rest)
        else:
            h.add_body(doc, plain(stripped))
        i += 1

    props = doc.core_properties
    props.author = "sssiqwel"
    props.last_modified_by = "sssiqwel"
    props.title = "Отчёт по лабораторной работе. Метод «большого взрыва»"
    props.comments = ""

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
